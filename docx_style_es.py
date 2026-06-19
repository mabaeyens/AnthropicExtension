# -*- coding: utf-8 -*-
"""
Shared styling/helpers for the editable Word (.docx) versions of the Spanish
documents for the Anthropic AI Assistant Qlik Sense extension.

Look & feel: Qlik-green professional, mirroring the PDF build. Pure python-docx.
Used by build_userguide_es_docx.py and build_architecture_es_docx.py.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# --- Palette -----------------------------------------------------------------
QLIK_GREEN = RGBColor(0x00, 0x98, 0x45)
QLIK_GREEN_DARK = RGBColor(0x00, 0x7A, 0x38)
DARK = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WARN_TEXT = RGBColor(0x7A, 0x5C, 0x00)

GREEN_HEX = '009845'
GREEN_DARK_HEX = '007A38'
ZEBRA_HEX = 'F7FAF8'
WARN_BG_HEX = 'FFF8E6'
WARN_BORDER_HEX = 'E0A800'
RULE_HEX = 'D7DCD9'
LIGHT_GREY_HEX = 'F2F4F3'

FONT = 'Calibri'
FONT_MONO = 'Consolas'


# --- low-level XML helpers ---------------------------------------------------
def _shade(el, hex_fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    el.append(shd)


def cell_bg(cell, hex_fill):
    _shade(cell._tc.get_or_add_tcPr(), hex_fill)


def _set_border(el_pr, edges, hex_color, sz='8', val='single'):
    borders = OxmlElement('w:tcBorders') if el_pr.tag.endswith('tcPr') else OxmlElement('w:pBdr')
    for edge in edges:
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), val)
        e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '4')
        e.set(qn('w:color'), hex_color)
        borders.append(e)
    el_pr.append(borders)


def cell_left_border(cell, hex_color, sz='24'):
    _set_border(cell._tc.get_or_add_tcPr(), ['left'], hex_color, sz=sz)


def cell_box(cell, hex_color, sz='8'):
    _set_border(cell._tc.get_or_add_tcPr(), ['top', 'bottom', 'left', 'right'], hex_color, sz=sz)


def para_bottom_border(paragraph, hex_color, sz='12'):
    _set_border(paragraph._p.get_or_add_pPr(), ['bottom'], hex_color, sz=sz)


def para_top_border(paragraph, hex_color, sz='18'):
    _set_border(paragraph._p.get_or_add_pPr(), ['top'], hex_color, sz=sz)


def add_page_field(paragraph):
    """Insert a dynamic PAGE field into a paragraph run."""
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instr); run._r.append(fldChar2)
    return run


# --- run / paragraph helpers -------------------------------------------------
def _run(p, text, bold=False, italic=False, color=DARK, size=None, mono=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    r.font.name = FONT_MONO if mono else FONT
    if size:
        r.font.size = Pt(size)
    return r


def add_runs(p, segments):
    """segments: list of dicts with text + optional bold/italic/color/mono/size."""
    for seg in segments:
        _run(p, seg['text'], bold=seg.get('bold', False), italic=seg.get('italic', False),
             color=seg.get('color', DARK), size=seg.get('size'), mono=seg.get('mono', False))


import re as _re
_TOK = _re.compile(r'(\*\*.+?\*\*|`.+?`)')


def parse(text):
    """Tiny markup: **bold** and `mono` -> run-segments for add_runs()."""
    segs = []
    for part in _TOK.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            segs.append({'text': part[2:-2], 'bold': True})
        elif part.startswith('`') and part.endswith('`'):
            segs.append({'text': part[1:-1], 'mono': True})
        else:
            segs.append({'text': part})
    return segs


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    _run(p, text, bold=True, color=QLIK_GREEN_DARK, size=16)
    para_bottom_border(p, RULE_HEX)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    _run(p, text, bold=True, color=DARK, size=12.5)
    return p


def h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    _run(p, text, bold=True, color=GREY, size=10.5)
    return p


def _apply_defaults(segs, size, color):
    for s in segs:
        s.setdefault('size', size)
        s.setdefault('color', color)
    return segs


def body(doc, segments_or_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    segs = parse(segments_or_text) if isinstance(segments_or_text, str) else segments_or_text
    add_runs(p, _apply_defaults(segs, 10.5, DARK))
    return p


def lead(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    _run(p, text, color=GREY, size=11)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    _run(p, text, italic=True, color=GREY, size=9)
    return p


def code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.left_indent = Cm(0.3)
    cell = None
    _run(p, text, color=DARK, size=9, mono=True)
    _shade(p._p.get_or_add_pPr(), LIGHT_GREY_HEX)
    return p


def bullet(doc, segments_or_text, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6 + 0.5 * level)
    p.paragraph_format.space_after = Pt(2)
    glyph = u'▪  ' if level == 0 else u'–  '
    gcolor = QLIK_GREEN if level == 0 else GREY
    _run(p, glyph, color=gcolor, size=10.5, bold=(level == 0))
    default_color = DARK if level == 0 else GREY
    segs = parse(segments_or_text) if isinstance(segments_or_text, str) else segments_or_text
    add_runs(p, _apply_defaults(segs, 10.5, default_color))
    return p


def bullets(doc, items):
    """items: list of (level, str-or-segments)."""
    for level, content in items:
        bullet(doc, content, level=level)


# --- tables ------------------------------------------------------------------
def spec_table(doc, headers, rows, widths_cm=None, mono_cols=None):
    mono_cols = mono_cols or set()
    n = len(headers)
    table = doc.add_table(rows=1, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = 'Table Grid'
    table.autofit = False
    hdr = table.rows[0].cells
    for i, htext in enumerate(headers):
        cell_bg(hdr[i], GREEN_HEX)
        para = hdr[i].paragraphs[0]
        para.paragraph_format.space_after = Pt(1)
        _run(para, htext, bold=True, color=WHITE, size=9.5)
    for r in rows:
        cells = table.add_row().cells
        for i, val in enumerate(r):
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(1)
            _run(para, str(val), size=9, mono=(i in mono_cols))
    # zebra striping on body rows
    for ri in range(1, len(rows) + 1):
        if ri % 2 == 0:
            for c in table.rows[ri].cells:
                cell_bg(c, ZEBRA_HEX)
    if widths_cm:
        for ri, row in enumerate(table.rows):
            for ci, c in enumerate(row.cells):
                c.width = Cm(widths_cm[ci])
    return table


def disclaimer_box(doc, lines, title=u'Aviso — demostración experimental'):
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.rows[0].cells[0]
    cell_bg(cell, WARN_BG_HEX)
    cell_box(cell, WARN_BORDER_HEX, sz='6')
    cell_left_border(cell, WARN_BORDER_HEX, sz='24')
    tp = cell.paragraphs[0]
    tp.paragraph_format.space_after = Pt(4)
    _run(tp, title, bold=True, color=WARN_TEXT, size=10)
    for ln in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        add_runs(p, ln)
    return table


# Disclaimer content as run-segments (bold where needed).
DISCLAIMER_LINES = [
    [
        {'text': u'Esta extensión es un activo de '},
        {'text': u'demostración experimental (v0.3.3, compilación 24)', 'bold': True},
        {'text': u', no un producto. '},
        {'text': u'No', 'bold': True},
        {'text': u' está reforzada para producción y '},
        {'text': u'no', 'bold': True},
        {'text': u' es una oferta de Qlik ni una integración soportada.'},
    ],
    [
        {'text': u'Ni Qlik ni el autor asumen responsabilidad alguna', 'bold': True},
        {'text': u' por cualquier incidencia, exposición de datos, coste o daño derivado de su '
                 u'uso en cualquier entorno de cliente o producción. Úselo bajo su propia '
                 u'responsabilidad.'},
    ],
    [
        {'text': u'Al hacer una pregunta, los datos del gráfico seleccionado salen del entorno hacia '
                 u'un LLM externo (Anthropic). No lo utilice con datos sensibles, regulados o personales '
                 u'sin autorización expresa.'},
    ],
]


# --- document scaffolding ----------------------------------------------------
def new_document():
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
    return doc


def add_footer(doc, foot_text):
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para_top_border(p, RULE_HEX, sz='6')
    _run(p, foot_text + u'      ', color=GREY, size=8)
    _run(p, u'Página ', color=GREY, size=8)
    pr = add_page_field(p)
    pr.font.size = Pt(8)
    pr.font.color.rgb = GREY
    pr.font.name = FONT


def cover(doc, doc_kind, doc_title, subtitle):
    # kind band (green shaded single-cell table)
    band = doc.add_table(rows=1, cols=1)
    bc = band.rows[0].cells[0]
    cell_bg(bc, GREEN_HEX)
    bp = bc.paragraphs[0]
    _run(bp, doc_kind.upper(), bold=True, color=WHITE, size=10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    tp = doc.add_paragraph()
    tp.paragraph_format.space_after = Pt(3)
    _run(tp, doc_title, bold=True, color=DARK, size=26)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(8)
    _run(sp, subtitle, color=GREY, size=13)
    para_bottom_border(sp, RULE_HEX)

    meta = [
        (u'Producto', u'Anthropic AI Assistant para Qlik Sense'),
        (u'Versión', u'v0.3.3 · compilación 24'),
        (u'Autor', u'mabaeyens'),
        (u'Fecha', u'19 de junio de 2026'),
        (u'Idioma', u'Español'),
    ]
    t = doc.add_table(rows=0, cols=2)
    for k, v in meta:
        cells = t.add_row().cells
        _run(cells[0].paragraphs[0], k, bold=True, size=10)
        _run(cells[1].paragraphs[0], v, size=10)
        cells[0].width = Cm(4)
        cells[1].width = Cm(13)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    disclaimer_box(doc, DISCLAIMER_LINES)
